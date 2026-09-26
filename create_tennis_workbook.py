#!/usr/bin/env python3
"""Create the initial operational tennis update workbook.

This is only an initializer for the preview branch. Once
LSC_Tennis_Update_System.xlsx exists, the script leaves it untouched so future
manual edits to the operational workbook are never overwritten.
"""

from datetime import date
from pathlib import Path

from openpyxl import Workbook
import pycountry
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUTPUT = Path("LSC_Tennis_Update_System.xlsx")

if OUTPUT.exists():
    print(f"{OUTPUT} already exists; leaving it untouched.")
    raise SystemExit(0)

wb = Workbook()
ws_readme = wb.active
ws_readme.title = "README"
ws_baseline = wb.create_sheet("Baseline")
ws_updates = wb.create_sheet("Updates")
ws_players = wb.create_sheet("Players")
ws_countries = wb.create_sheet("Countries")
ws_lists = wb.create_sheet("Lists")
ws_examples = wb.create_sheet("Examples")

navy = "07141C"
green = "3E8B68"
green_bright = "7FBEA1"
gold = "C49A47"
cream = "F1EBDD"
muted = "D8E0DD"
line = "29413F"
white = "FFFFFF"

thin = Side(style="thin", color=line)
header_fill = PatternFill("solid", fgColor=navy)
green_fill = PatternFill("solid", fgColor=green)
gold_fill = PatternFill("solid", fgColor=gold)

# README
ws_readme["A1"] = "LSC Tennis Update System"
ws_readme["A1"].font = Font(size=18, bold=True, color=cream)
ws_readme["A3"] = "Purpose"
ws_readme["A3"].font = Font(bold=True, color=green_bright)
ws_readme["A4"] = (
    "Operational workbook for the MEN'S and WOMEN'S Open Era lineal tennis titles. "
    "It does not contain the full historical reconstruction. Historical research/audit "
    "masters remain archived separately. This workbook holds the frozen current baseline "
    "plus future eligible match updates only."
)
ws_readme["A4"].alignment = Alignment(wrap_text=True, vertical="top")
ws_readme["A6"] = "Player representation / flags"
ws_readme["A6"].font = Font(bold=True, color=green_bright)
ws_readme["A7"] = ("Players contains one row per player. Choose a Country from the Countries lookup. "
                    "The builder derives the Unicode flag automatically. Use Flag Override only for "
                    "historical/special representations where a modern flag would be misleading.")
ws_readme["A7"].alignment = Alignment(wrap_text=True, vertical="top")
ws_readme["A8"] = "How to use"
ws_readme["A8"].font = Font(bold=True, color=green_bright)
rules = [
    "Do not edit the Baseline unless a new publication baseline is deliberately frozen.",
    "For each future holder match, append one row to Updates.",
    "Scheduled/Awaiting result rows must leave Result Type and Played? blank.",
    "For a completed normal match or played retirement/default, set Played? = Yes.",
    "For a walkover or other non-played result, set Played? = No. It must not transfer or count as a defence.",
    "Holder win = successful defence. Challenger win = title transfer. No result = no lineage change.",
    "Surface must be Hard, Clay, Grass or Carpet for played matches.",
    "Completed matches require a source URL.",
    "Do not paste historical audit rows here. The workbook begins from the frozen publication checkpoint.",
]
for i, rule in enumerate(rules, start=9):
    ws_readme.cell(i, 1, i - 8)
    ws_readme.cell(i, 2, rule)
    ws_readme.cell(i, 2).alignment = Alignment(wrap_text=True)
ws_readme.column_dimensions["A"].width = 12
ws_readme.column_dimensions["B"].width = 100

# Baseline
baseline_headers = [
    "Lineage", "Tour", "Baseline Holder", "Baseline Since", "Won From",
    "Tournament", "Round", "Surface", "Acquisition Score", "Current Defences",
    "Transfers", "Total Defences", "Played Title Matches", "Checkpoint Date", "Notes",
]
ws_baseline.append(baseline_headers)
ws_baseline.append([
    "Men's Singles", "ATP", "Jannik Sinner", date(2026, 7, 12), "Alexander Zverev",
    "Wimbledon", "F", "Grass", "6–7(7) 7–6(2) 6–3 6–4", 0,
    1610, 4072, 5682, date(2026, 9, 18),
    "Frozen publication baseline. No successful defence after acquisition through checkpoint.",
])
ws_baseline.append([
    "Women's Singles", "WTA", "Elena Rybakina", date(2026, 9, 11), "Coco Gauff",
    "US Open", "SF", "Hard", "3–6, 6–4, 6–4", 1,
    1087, 3381, 4468, date(2026, 9, 12),
    "Frozen publication baseline. One successful defence: 12 Sep 2026 US Open F vs Aryna Sabalenka.",
])

# Updates
update_headers = [
    "Lineage", "Match Date", "Opponent", "Tournament", "Round", "Surface",
    "Status", "Played?", "Result Type", "Result / Score", "Source URL", "Notes",
]
ws_updates.append(update_headers)
ws_updates["N1"] = "LIVE SHEET"
ws_updates["N1"].font = Font(bold=True, color=green_bright)
ws_updates["N2"] = "Append only real holder matches here. See Examples for test patterns."
ws_updates["N2"].alignment = Alignment(wrap_text=True)

# Players
player_headers = ["Player", "Country", "Country Code", "Flag Override", "Notes"]
ws_players.append(player_headers)
ws_players.append(["Jannik Sinner", "Italy", "ITA", "", "Current men's baseline holder."])
ws_players.append(["Elena Rybakina", "Kazakhstan", "KAZ", "", "Current women's baseline holder."])

# Countries — standard ISO country lookup with generated Unicode flags.
country_rows = []
for country in sorted(pycountry.countries, key=lambda c: c.name):
    alpha2 = country.alpha_2
    flag = "".join(chr(127397 + ord(ch)) for ch in alpha2)
    country_rows.append([country.name, country.alpha_3, alpha2, flag])

ws_countries.append(["Country", "ISO3", "ISO2", "Flag"])
for row in country_rows:
    ws_countries.append(row)

# Lists
list_columns = {
    "A": ("Lineage", ["Men's Singles", "Women's Singles"]),
    "B": ("Surface", ["Hard", "Clay", "Grass", "Carpet"]),
    "C": ("Status", ["Scheduled", "Awaiting result", "Completed", "Postponed", "Cancelled"]),
    "D": ("Played?", ["Yes", "No"]),
    "E": ("Result Type", ["Holder win", "Challenger win", "No result"]),
    "F": ("Round", ["R128", "R64", "R32", "R16", "QF", "SF", "F", "RR", "Other"]),
}
for col, (head, vals) in list_columns.items():
    ws_lists[f"{col}1"] = head
    for r, value in enumerate(vals, start=2):
        ws_lists[f"{col}{r}"] = value

# Examples
ws_examples.append(update_headers)
ws_examples.append([
    "Men's Singles", date(2026, 10, 1), "Example Opponent", "Example Event",
    "R16", "Hard", "Completed", "Yes", "Holder win", "6–4 6–3",
    "https://example.com", "Demonstrates a successful defence.",
])
ws_examples.append([
    "Men's Singles", date(2026, 10, 8), "Example Challenger", "Example Event",
    "QF", "Hard", "Completed", "Yes", "Challenger win", "4–6 6–3 6–2",
    "https://example.com", "Demonstrates a title transfer.",
])
ws_examples.append([
    "Women's Singles", date(2026, 10, 15), "Example Opponent", "Example Event",
    "SF", "Clay", "Completed", "No", "No result", "Walkover",
    "https://example.com", "Demonstrates a non-played match; no lineage change.",
])

# Shared formatting
for ws in [ws_baseline, ws_updates, ws_players, ws_countries, ws_lists, ws_examples]:
    for cell in ws[1]:
        cell.fill = green_fill
        cell.font = Font(bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    ws.freeze_panes = "A2"

for ws in [ws_baseline, ws_updates, ws_players, ws_examples]:
    for col in range(1, 16):
        ws.column_dimensions[chr(64 + col)].width = 16
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["I"].width = 26
    ws.column_dimensions["J"].width = 18
    ws.column_dimensions["K"].width = 34
    ws.column_dimensions["L"].width = 50
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

for cell in ["D2", "N2", "D3", "N3"]:
    ws_baseline[cell].number_format = "yyyy-mm-dd"
for row in range(2, 301):
    ws_updates.cell(row, 2).number_format = "yyyy-mm-dd"
for row in range(2, 20):
    ws_examples.cell(row, 2).number_format = "yyyy-mm-dd"

# Validations on the live Updates sheet.
def add_list_validation(cell_range, formula):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws_updates.add_data_validation(dv)
    dv.add(cell_range)

add_list_validation("A2:A500", "Lists!$A$2:$A$3")
dv_country = DataValidation(type="list", formula1=f"Countries!$A$2:$A{len(country_rows)+1{'}'}", allow_blank=True)
ws_players.add_data_validation(dv_country)
dv_country.add("B2:B1000")
add_list_validation("E2:E500", "Lists!$F$2:$F$10")
add_list_validation("F2:F500", "Lists!$B$2:$B$5")
add_list_validation("G2:G500", "Lists!$C$2:$C$6")
add_list_validation("H2:H500", "Lists!$D$2:$D$3")
add_list_validation("I2:I500", "Lists!$E$2:$E$4")

# Visual hints.
for row in range(2, 501):
    ws_updates.cell(row, 1).fill = PatternFill("solid", fgColor="F4F8F6")
ws_updates["N1"].fill = gold_fill

wb.save(OUTPUT)
print(f"Created {OUTPUT}")
