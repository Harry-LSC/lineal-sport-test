#!/usr/bin/env python3
"""Seed three synthetic men's tennis updates for pipeline regression testing."""

from datetime import date
from pathlib import Path
from openpyxl import load_workbook

PATH = Path("LSC_Tennis_Update_System.xlsx")
wb = load_workbook(PATH)
ws = wb["Updates"]

# Remove prior synthetic rows from this test branch so reruns are deterministic.
for row in range(ws.max_row, 1, -1):
    if ws.cell(row, 12).value and str(ws.cell(row, 12).value).startswith("PIPELINE TEST"):
        ws.delete_rows(row, 1)

rows = [
    [
        "Men's Singles", date(2026, 9, 19), "Test Opponent A", "Pipeline Test Event",
        "QF", "Hard", "Completed", "Yes", "Holder win", "6–4 6–4",
        "https://example.com/test-defence-1", "PIPELINE TEST: holder defence",
    ],
    [
        "Men's Singles", date(2026, 9, 20), "Test Challenger", "Pipeline Test Event",
        "SF", "Hard", "Completed", "Yes", "Challenger win", "4–6 6–3 6–2",
        "https://example.com/test-transfer", "PIPELINE TEST: title transfer",
    ],
    [
        "Men's Singles", date(2026, 9, 21), "Test Opponent B", "Pipeline Test Event",
        "F", "Hard", "Completed", "Yes", "Holder win", "7–5 6–3",
        "https://example.com/test-defence-2", "PIPELINE TEST: new holder defence",
    ],
]

for row in rows:
    ws.append(row)

wb.save(PATH)
print("Seeded 3 synthetic updates.")
