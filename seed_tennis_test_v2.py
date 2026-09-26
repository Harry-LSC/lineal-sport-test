#!/usr/bin/env python3
from datetime import date
from openpyxl import load_workbook

path="LSC_Tennis_Update_System.xlsx"
wb=load_workbook(path)
players=wb["Players"]
updates=wb["Updates"]

# Add synthetic player representation once.
names={players.cell(r,1).value for r in range(2,players.max_row+1)}
if "Test Challenger" not in names:
    players.append(["Test Challenger","France","FRA","","PIPELINE TEST synthetic player"])

# Clear prior synthetic update rows if any.
for r in range(updates.max_row,1,-1):
    note=updates.cell(r,12).value
    if note and str(note).startswith("PIPELINE TEST"):
        updates.delete_rows(r,1)

rows=[
["Men's Singles",date(2026,9,19),"Test Opponent A","Pipeline Test Event","QF","Hard","Completed","Yes","Holder win","6–4 6–4","https://example.com/test-defence-1","PIPELINE TEST holder defence"],
["Men's Singles",date(2026,9,20),"Test Challenger","Pipeline Test Event","SF","Hard","Completed","Yes","Challenger win","4–6 6–3 6–2","https://example.com/test-transfer","PIPELINE TEST title transfer"],
["Men's Singles",date(2026,9,21),"Test Opponent B","Pipeline Test Event","F","Hard","Completed","Yes","Holder win","7–5 6–3","https://example.com/test-defence-2","PIPELINE TEST new holder defence"],
]
for row in rows:
    updates.append(row)

wb.save(path)
print("Synthetic player + three matches seeded.")
