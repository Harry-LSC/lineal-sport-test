#!/usr/bin/env python3
"""
Lineal Sport Championship — Tennis Excel -> JSON builder

Reads:
    LSC_Tennis_Update_System.xlsx

Writes:
    data/tennis-data.json

The workbook is intentionally operational, not archival. It contains a frozen
baseline for Men's Singles and Women's Singles plus future holder-match updates.
Full historical reconstruction/audit masters remain offline.
"""

from __future__ import annotations

import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

WORKBOOK = Path("LSC_Tennis_Update_System.xlsx")
OUTPUT = Path("data/tennis-data.json")

VALID_LINEAGES = {"Men's Singles", "Women's Singles"}
VALID_TOURS = {"ATP", "WTA"}
VALID_SURFACES = {"Hard", "Clay", "Grass", "Carpet"}
VALID_STATUSES = {"Scheduled", "Awaiting result", "Completed", "Postponed", "Cancelled"}
VALID_PLAYED = {"Yes", "No"}
VALID_RESULTS = {"Holder win", "Challenger win", "No result"}


def clean(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return value


def iso_date(value):
    if value in ("", None):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raise ValueError(f"Expected an Excel date, got {value!r}")


def fail(message):
    print(f"ERROR: {message}", file=sys.stderr)
    raise SystemExit(1)


def header_map(ws, required):
    headers = [clean(c.value) for c in ws[1]]
    missing = [h for h in required if h not in headers]
    if missing:
        fail(f"{ws.title} sheet is missing columns: {', '.join(missing)}")
    return {name: headers.index(name) for name in required}


def read_baseline(ws):
    required = [
        "Lineage", "Tour", "Baseline Holder", "Baseline Since", "Won From",
        "Tournament", "Round", "Surface", "Acquisition Score",
        "Current Defences", "Transfers", "Total Defences",
        "Played Title Matches", "Checkpoint Date", "Notes",
    ]
    idx = header_map(ws, required)
    result = {}

    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        lineage = clean(row[idx["Lineage"]])
        if not lineage:
            continue
        if lineage not in VALID_LINEAGES:
            fail(f"Baseline row {r}: unknown lineage '{lineage}'.")
        if lineage in result:
            fail(f"Baseline row {r}: duplicate lineage '{lineage}'.")

        tour = clean(row[idx["Tour"]])
        if tour not in VALID_TOURS:
            fail(f"Baseline row {r}: invalid Tour '{tour}'.")

        surface = clean(row[idx["Surface"]])
        if surface not in VALID_SURFACES:
            fail(f"Baseline row {r}: invalid acquisition Surface '{surface}'.")

        holder = clean(row[idx["Baseline Holder"]])
        if not holder:
            fail(f"Baseline row {r}: Baseline Holder is blank.")

        baseline_since = iso_date(row[idx["Baseline Since"]])
        checkpoint_date = iso_date(row[idx["Checkpoint Date"]])
        if not baseline_since or not checkpoint_date:
            fail(f"Baseline row {r}: Baseline Since and Checkpoint Date are required.")
        if checkpoint_date < baseline_since:
            fail(f"Baseline row {r}: Checkpoint Date precedes Baseline Since.")

        result[lineage] = {
            "lineage": lineage,
            "tour": tour,
            "current_holder": holder,
            "current_since": baseline_since,
            "won_from": clean(row[idx["Won From"]]) or None,
            "acquisition": {
                "tournament": clean(row[idx["Tournament"]]) or None,
                "round": clean(row[idx["Round"]]) or None,
                "surface": surface,
                "score": str(clean(row[idx["Acquisition Score"]])) or None,
            },
            "current_defences": int(row[idx["Current Defences"]] or 0),
            "transfers": int(row[idx["Transfers"]] or 0),
            "total_defences": int(row[idx["Total Defences"]] or 0),
            "played_title_matches": int(row[idx["Played Title Matches"]] or 0),
            "checkpoint_date": checkpoint_date,
            "notes": clean(row[idx["Notes"]]) or None,
        }

    if set(result) != VALID_LINEAGES:
        missing = sorted(VALID_LINEAGES - set(result))
        fail("Baseline must contain exactly the two tennis lineages; missing: " + ", ".join(missing))

    return result


def country_flag_from_alpha2(alpha2):
    if not alpha2 or len(alpha2) != 2:
        return None
    return "".join(chr(127397 + ord(ch)) for ch in alpha2.upper())


def read_players(ws_players, ws_countries):
    country_headers = [clean(c.value) for c in ws_countries[1]]
    for req in ("Country", "ISO3", "ISO2", "Flag"):
        if req not in country_headers:
            fail(f"Countries sheet is missing column: {req}")
    cidx = {h: country_headers.index(h) for h in country_headers}
    countries = {}
    for row in ws_countries.iter_rows(min_row=2, values_only=True):
        name = clean(row[cidx["Country"]])
        if not name:
            continue
        countries[name] = {
            "country": name,
            "country_code": clean(row[cidx["ISO3"]]) or None,
            "iso2": clean(row[cidx["ISO2"]]) or None,
            "flag": clean(row[cidx["Flag"]]) or country_flag_from_alpha2(clean(row[cidx["ISO2"]])),
        }

    required = ["Player", "Country", "Country Code", "Flag Override", "Notes"]
    idx = header_map(ws_players, required)
    players = {}
    for r, row in enumerate(ws_players.iter_rows(min_row=2, values_only=True), start=2):
        player = clean(row[idx["Player"]])
        if not player:
            continue
        if player in players:
            fail(f"Players row {r}: duplicate player '{player}'.")
        country = clean(row[idx["Country"]])
        if not country:
            fail(f"Players row {r}: Country is blank for '{player}'.")
        if country not in countries:
            fail(f"Players row {r}: unknown Country '{country}' for '{player}'.")
        country_code = clean(row[idx["Country Code"]]) or countries[country]["country_code"]
        flag_override = clean(row[idx["Flag Override"]])
        players[player] = {
            "player": player,
            "country": country,
            "country_code": country_code,
            "flag": flag_override or countries[country]["flag"],
            "flag_override": flag_override or None,
            "notes": clean(row[idx["Notes"]]) or None,
        }
    return players


def read_updates(ws, baseline):
    required = [
        "Lineage", "Match Date", "Opponent", "Tournament", "Round", "Surface",
        "Status", "Played?", "Result Type", "Result / Score", "Source URL", "Notes",
    ]
    idx = header_map(ws, required)
    events = []

    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        lineage = clean(row[idx["Lineage"]])
        match_date_raw = row[idx["Match Date"]]
        opponent = clean(row[idx["Opponent"]])
        tournament = clean(row[idx["Tournament"]])
        status = clean(row[idx["Status"]])

        if not any([lineage, match_date_raw, opponent, tournament, status]):
            continue

        if lineage not in baseline:
            fail(f"Updates row {r}: unknown Lineage '{lineage}'.")
        if not match_date_raw:
            fail(f"Updates row {r}: Match Date is blank.")
        if not opponent:
            fail(f"Updates row {r}: Opponent is blank.")
        if not tournament:
            fail(f"Updates row {r}: Tournament is blank.")
        if status not in VALID_STATUSES:
            fail(f"Updates row {r}: invalid Status '{status}'.")

        match_date = iso_date(match_date_raw)
        if match_date <= baseline[lineage]["checkpoint_date"]:
            fail(
                f"Updates row {r}: Match Date {match_date} is not after the frozen "
                f"checkpoint {baseline[lineage]['checkpoint_date']} for {lineage}."
            )

        round_name = clean(row[idx["Round"]])
        surface = clean(row[idx["Surface"]])
        played = clean(row[idx["Played?"]])
        result_type = clean(row[idx["Result Type"]])
        score = clean(row[idx["Result / Score"]])
        source_url = clean(row[idx["Source URL"]])

        if surface and surface not in VALID_SURFACES:
            fail(f"Updates row {r}: invalid Surface '{surface}'.")

        if status == "Completed":
            if played not in VALID_PLAYED:
                fail(f"Updates row {r}: Completed match requires Played? = Yes or No.")
            if not source_url:
                fail(f"Updates row {r}: Completed match requires a Source URL.")

            if played == "Yes":
                if result_type not in {"Holder win", "Challenger win"}:
                    fail(
                        f"Updates row {r}: a played Completed match requires Result Type "
                        "Holder win or Challenger win."
                    )
                if not surface:
                    fail(f"Updates row {r}: a played Completed match requires Surface.")
                if not score:
                    fail(f"Updates row {r}: a played Completed match requires Result / Score.")
            else:
                if result_type != "No result":
                    fail(
                        f"Updates row {r}: Played? = No requires Result Type = No result "
                        "(e.g. walkover)."
                    )
        else:
            if played or result_type or score:
                fail(
                    f"Updates row {r}: Played?, Result Type and Result / Score must remain "
                    f"blank while Status is {status}."
                )

        events.append({
            "row": r,
            "lineage": lineage,
            "match_date": match_date,
            "opponent": opponent,
            "tournament": tournament,
            "round": round_name or None,
            "surface": surface or None,
            "status": status,
            "played": played or None,
            "result_type": result_type or None,
            "result_score": str(score) if score != "" else None,
            "source_url": source_url or None,
            "notes": clean(row[idx["Notes"]]) or None,
        })

    last_date = {lineage: baseline[lineage]["checkpoint_date"] for lineage in baseline}
    for event in events:
        lineage = event["lineage"]
        if event["match_date"] < last_date[lineage]:
            fail(
                f"Updates row {event['row']}: chronology goes backwards for {lineage} "
                f"({event['match_date']} < {last_date[lineage]})."
            )
        last_date[lineage] = event["match_date"]

    return events


def build_state(baseline, events, players):
    for lineage, state in baseline.items():
        holder = state["current_holder"]
        if holder not in players:
            fail(f"Baseline current holder '{holder}' ({lineage}) is missing from Players sheet.")

    states = {name: dict(values) for name, values in baseline.items()}
    for state in states.values():
        state["acquisition"] = dict(state["acquisition"])

    lineage_updates = []
    pending = {name: [] for name in baseline}

    for event in events:
        lineage = event["lineage"]
        state = states[lineage]
        status = event["status"]

        if status == "Completed":
            holder_before = state["current_holder"]

            if event["played"] == "No":
                lineage_updates.append({
                    **{k: v for k, v in event.items() if k != "row"},
                    "holder_before": holder_before,
                    "holder_after": holder_before,
                    "outcome": "No result",
                    "counts_as_title_match": False,
                })
                continue

            state["played_title_matches"] += 1

            if event["result_type"] == "Holder win":
                outcome = "Defence"
                holder_after = holder_before
                state["current_defences"] += 1
                state["total_defences"] += 1

            elif event["result_type"] == "Challenger win":
                outcome = "Transfer"
                holder_after = event["opponent"]
                if holder_after not in players:
                    fail(
                        f"New holder '{holder_after}' is missing from Players sheet. "
                        "Add the player and representation before publishing."
                    )
                state["current_holder"] = holder_after
                state["current_since"] = event["match_date"]
                state["won_from"] = holder_before
                state["current_defences"] = 0
                state["transfers"] += 1
                state["acquisition"] = {
                    "tournament": event["tournament"],
                    "round": event["round"],
                    "surface": event["surface"],
                    "score": event["result_score"],
                }
            else:
                fail(f"Internal error: unexpected result type {event['result_type']!r}")

            lineage_updates.append({
                **{k: v for k, v in event.items() if k != "row"},
                "holder_before": holder_before,
                "holder_after": holder_after,
                "outcome": outcome,
                "counts_as_title_match": True,
            })

        elif status in {"Scheduled", "Awaiting result"}:
            pending[lineage].append(event)

    lineages = []
    for lineage, state in states.items():
        candidates = sorted(pending[lineage], key=lambda e: (e["match_date"], e["row"]))
        next_event = candidates[0] if candidates else None

        lineages.append({
            **state,
            "next_match": ({
                "date": next_event["match_date"],
                "opponent": next_event["opponent"],
                "tournament": next_event["tournament"],
                "round": next_event["round"],
                "surface": next_event["surface"],
                "status": next_event["status"],
                "source_url": next_event["source_url"],
            } if next_event else None),
        })

    return {
        "schema_version": 1,
        "system": "LSC Tennis Update System",
        "lineages": lineages,
        "lineage_updates": lineage_updates,
        "players": list(players.values()),
    }


def main():
    if not WORKBOOK.exists():
        fail(f"Workbook not found: {WORKBOOK}")

    wb = load_workbook(WORKBOOK, data_only=False, read_only=True)
    for sheet in ("Baseline", "Updates", "Players", "Countries"):
        if sheet not in wb.sheetnames:
            fail(f"Workbook is missing required sheet: {sheet}")

    baseline = read_baseline(wb["Baseline"])
    players = read_players(wb["Players"], wb["Countries"])
    events = read_updates(wb["Updates"], baseline)
    output = build_state(baseline, events, players)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print(
        f"Built {OUTPUT}: {len(output['lineages'])} lineages, "
        f"{len(output['lineage_updates'])} completed/non-played update event(s)."
    )


if __name__ == "__main__":
    main()
