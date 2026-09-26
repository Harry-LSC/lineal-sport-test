#!/usr/bin/env python3
"""Add/maintain Players and Countries sheets in the operational tennis workbook."""

from pathlib import Path
import pycountry
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

PATH = Path("LSC_Tennis_Update_System.xlsx")
if not PATH.exists():
    raise SystemExit("Workbook missing; run create_tennis_workbook.py first.")

wb = load_workbook(PATH)

green = "3E8B68"
white = "FFFFFF"
header_fill = PatternFill("solid", fgColor=green)

if "Countries" in wb.sheetnames:
    del wb["Countries"]
ws_countries = wb.create_sheet("Countries", 4)
ws_countries.append(["Country", "ISO3", "ISO2", "Flag"])
country_rows = []
for country in sorted(pycountry.countries, key=lambda c: c.name):
    alpha2 = country.alpha_2
    flag = "".join(chr(127397 + ord(ch)) for ch in alpha2)
    country_rows.append([country.name, country.alpha_3, alpha2, flag])
for row in country_rows:
    ws_countries.append(row)

if "Players" not in wb.sheetnames:
    ws_players = wb.create_sheet("Players", 3)
    ws_players.append(["Player", "Country", "Country Code", "Flag Override", "Notes"])
    ws_players.append(["Jannik Sinner", "Italy", "ITA", "", "Current men's baseline holder."])
    ws_players.append(["Elena Rybakina", "Kazakhstan", "KAZ", "", "Current women's baseline holder."])
else:
    ws_players = wb["Players"]

for ws in (ws_players, ws_countries):
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"

for col, width in {"A":26, "B":26, "C":15, "D":18, "E":55}.items():
    ws_players.column_dimensions[col].width = width
for col, width in {"A":34, "B":12, "C":10, "D":10}.items():
    ws_countries.column_dimensions[col].width = width

ws_players.data_validations.dataValidation = []
dv_country = DataValidation(
    type="list",
    formula1="Countries!$A$2:$A$" + str(len(country_rows)+1),
    allow_blank=True,
)
ws_players.add_data_validation(dv_country)
dv_country.add("B2:B1000")

ws_readme = wb["README"]
note = "Player representation / flags"
existing = [ws_readme.cell(r, 1).value for r in range(1, ws_readme.max_row + 1)]
if note not in existing:
    insert_at = 6
    ws_readme.insert_rows(insert_at, amount=2)
    ws_readme.cell(insert_at, 1, note)
    ws_readme.cell(insert_at, 1).font = Font(bold=True, color="7FBEA1")
    ws_readme.cell(insert_at + 1, 1,
        "Players contains one row per player. Choose a Country from the Countries lookup. "
        "The builder derives the Unicode flag automatically. Use Flag Override only for "
        "historical/special representations where a modern flag would be misleading."
    )
    ws_readme.cell(insert_at + 1, 1).alignment = Alignment(wrap_text=True, vertical="top")

wb.save(PATH)
print("Players/Countries lookup ready.")
